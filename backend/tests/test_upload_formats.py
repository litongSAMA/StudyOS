"""上传识别：常见文件类型格式全覆盖的端到端用例（真实文件，不做 mock）。

每个格式走的是同一条真实链路：multipart 上传 → 落盘 → 内容嗅探 → 抽取器
→ DDL/资料建议。用例同时断言「抽到的正文」与「回包的格式元信息」，
因为前端要靠后者展示格式徽标与不一致提示。
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest

helpers_dir = Path(__file__).resolve().parent


def upload(client, name: str, data: bytes, content_type: str = "application/octet-stream"):
    return client.post(
        "/api/materials/upload",
        files={"file": (name, data, content_type)},
    )


def body_of(client, name: str, data: bytes, content_type: str = "application/octet-stream"):
    r = upload(client, name, data, content_type)
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------------------
# 文档 / 表格 / 演示
# ---------------------------------------------------------------------------

def test_upload_docx_with_table_and_deadline(client):
    """Word：正文 + 表格都要进正文（DDL 常写在表格里）。"""
    from docx import Document

    doc = Document()
    doc.add_heading("第三章 内存管理", level=1)
    doc.add_paragraph("本章作业提交截止 2026-12-01 23:59，请提交到教学平台。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "章节"
    table.cell(0, 1).text = "考点"
    table.cell(1, 0).text = "3.1 分页"
    table.cell(1, 1).text = "页表与快表"
    buf = io.BytesIO()
    doc.save(buf)

    body = body_of(client, "操作系统第三章.docx", buf.getvalue())
    assert body["format"]["ext"] == ".docx"
    assert body["format"]["label"] == "Word 文档"
    assert "分页" in body["extracted_text_preview"]
    assert "2026-12-01" in body["extracted_text_preview"]
    assert body["text_chars"] > 0


def test_upload_pptx_notes(client):
    """PPT：正文、表格与演讲者备注都要被抽出来。"""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "进程调度算法"
    box = slide.shapes.add_textbox(0, 0, 400, 100)
    box.text_frame.text = "先来先服务 / 短作业优先"
    slide.notes_slide.notes_text_frame.text = "备注：期末必考计算题"
    buf = io.BytesIO()
    prs.save(buf)

    body = body_of(client, "第4讲 调度.pptx", buf.getvalue())
    assert body["format"]["ext"] == ".pptx"
    preview = body["extracted_text_preview"]
    assert "进程调度算法" in preview and "先来先服务" in preview
    assert "期末必考" in preview          # 备注也是考点来源


def test_upload_xlsx_sheets(client):
    """Excel：多工作表渲染成表格文本（成绩单/名单里的日期不能丢）。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(["学号", "姓名", "成绩"])
    ws.append(["2021001", "张三", 92])
    ws2 = wb.create_sheet("安排")
    ws2.append(["实验", "2026-11-20 14:00"])
    buf = io.BytesIO()
    wb.save(buf)

    body = body_of(client, "成绩单.xlsx", buf.getvalue())
    assert body["format"]["ext"] == ".xlsx"
    assert body["format"]["group"] == "sheet"
    preview = body["extracted_text_preview"]
    assert "成绩" in preview and "张三" in preview
    assert "2026-11-20" in preview


def test_upload_odt(client):
    """OpenDocument 文本：段落 + 表格（不能出现表格单元格重复输出）。"""
    body = body_of(client, "讲义.odt", _odt_bytes(
        "<text:h>第四章 文件系统</text:h>"
        "<text:p>实验报告 2026-12-05 前提交</text:p>"
        "<table:table table:name=\"T\"><table:table-row>"
        "<table:table-cell><text:p>索引</text:p></table:table-cell>"
        "<table:table-cell><text:p>考点</text:p></table:table-cell>"
        "</table:table-row></table:table>"
    ))
    assert body["format"]["ext"] == ".odt"
    preview = body["extracted_text_preview"]
    assert "文件系统" in preview
    assert preview.count("索引") == 1        # 表格单元格只输出一次


def test_upload_csv_rendered_as_table(client):
    """CSV 渲染成「表头 + 行号」的表格文本，而不是原样倒分隔符。"""
    body = body_of(client, "点名册.csv", "学号,姓名,提交日期\n1,张三,2026-11-11\n".encode("utf-8"))
    assert body["format"]["ext"] == ".csv"
    preview = body["extracted_text_preview"]
    assert "表头" in preview and "张三" in preview


def test_upload_gbk_csv(client):
    """GBK 编码的教务导出文件不能变成乱码。"""
    raw = "课程,教师\n高等数学,李四\n".encode("gbk")
    body = body_of(client, "课程名单.csv", raw)
    assert "高等数学" in body["extracted_text_preview"]
    assert "乱码" not in body["extracted_text_preview"]


# ---------------------------------------------------------------------------
# 电子书 / Notebook / 网页 / RTF / 字幕 / 代码 / 压缩包
# ---------------------------------------------------------------------------

def test_upload_epub(client):
    """EPUB：按 spine 顺序抽正文。"""
    body = body_of(client, "操作系统概念.epub", _epub_bytes([
        ("ch1.xhtml", "<html><body><h1>第一章 概述</h1><p>进程与线程</p></body></html>"),
        ("ch2.xhtml", "<html><body><h1>第二章 内存</h1><p>分页与分段</p></body></html>"),
    ]))
    assert body["format"]["ext"] == ".epub"
    preview = body["extracted_text_preview"]
    assert "第一章" in preview and "分页与分段" in preview


def test_upload_ipynb(client):
    """Notebook：Markdown 单元格与代码/输出都要，且不被当成普通 JSON。"""
    nb = {
        "cells": [
            {"cell_type": "markdown", "source": ["# 实验三 页面置换"]},
            {"cell_type": "code", "source": ["print('LRU')"],
             "outputs": [{"output_type": "stream", "text": ["LRU\n"]}]},
        ],
        "nbformat": 4,
    }
    body = body_of(client, "实验三.ipynb", json.dumps(nb, ensure_ascii=False).encode("utf-8"))
    assert body["format"]["ext"] == ".ipynb"
    assert body["format"]["group"] == "notebook"
    preview = body["extracted_text_preview"]
    assert "页面置换" in preview and "LRU" in preview


def test_upload_html(client):
    """网页：标题/段落/表格都要，脚本与样式必须丢弃。"""
    html = (
        "<html><head><title>课程通知</title><style>.x{color:red}</style></head><body>"
        "<h1>期末考试安排</h1><p>考试时间 2026-12-20 09:00</p>"
        "<script>alert('不该出现')</script>"
        "<table><tr><td>科目</td><td>时间</td></tr></table></body></html>"
    )
    body = body_of(client, "通知.html", html.encode("utf-8"))
    assert body["format"]["ext"] == ".html"
    preview = body["extracted_text_preview"]
    assert "期末考试安排" in preview
    assert "alert" not in preview and "color:red" not in preview
    assert "科目" in preview


def test_upload_rtf(client):
    """RTF：\\'hh 十六进制转义要按代码页还原成中文。"""
    rtf = (r"{\rtf1\ansi\ansicpg936\deff0{\fonttbl{\f0\fnil Arial;}}"
           r"\viewkind4\uc1\pard\f0\fs24 \'b8\'df\'b5\'c8\'ca\'fd\'d1\'a7\par "
           r"6\'d4\'c21\'c8\'d5\'bf\'bc\'ca\'d4\par}")
    body = body_of(client, "考试说明.rtf", rtf.encode("latin-1"))
    assert body["format"]["ext"] == ".rtf"
    assert "高等数学" in body["extracted_text_preview"]
    assert "6月1日" in body["extracted_text_preview"]


def test_upload_srt(client):
    """字幕：保留时间戳，便于把知识点定位到视频位置。"""
    srt = ("1\n00:00:10,000 --> 00:00:14,000\n第三章讲的是内存管理\n\n"
           "2\n00:01:20,000 --> 00:01:25,000\n考点是页面置换算法\n")
    body = body_of(client, "网课字幕.srt", srt.encode("utf-8"))
    assert body["format"]["ext"] == ".srt"
    preview = body["extracted_text_preview"]
    assert "[00:10]" in preview and "[01:20]" in preview
    assert "页面置换算法" in preview


def test_upload_code_file(client):
    """代码：给顶层定义摘要，方便「这段代码讲了什么」的提炼。"""
    code = "import os\n\n\ndef parse_schedule(path):\n    return path\n\n\nclass Timetable:\n    pass\n"
    body = body_of(client, "timetable.py", code.encode("utf-8"))
    assert body["format"]["ext"] == ".py"
    assert body["format"]["group"] == "code"
    preview = body["extracted_text_preview"]
    assert "顶层定义" in preview and "parse_schedule" in preview


def test_upload_zip_expands_supported_entries(client):
    """压缩包：展开包内可解析的文档与文本，合并成一份正文。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("包内 Word：实验报告 2026-12-10 前提交")
    buf = io.BytesIO()
    doc.save(buf)

    package = io.BytesIO()
    with zipfile.ZipFile(package, "w") as zf:
        zf.writestr("课程/readme.txt", "包内说明：考试范围见附件")
        zf.writestr("课程/实验报告.docx", buf.getvalue())
        zf.writestr("课程/幻灯.pptx", b"not really a pptx")      # 坏文件不能中断整包
    body = body_of(client, "资料包.zip", package.getvalue())
    assert body["format"]["ext"] == ".zip"
    preview = body["extracted_text_preview"]
    assert "包内说明" in preview
    assert "实验报告" in preview
    assert any("跳过" in w or "压缩包" in w for w in body["extract"]["warnings"])


def test_upload_tar_gz_expands_entries(client):
    """tar.gz（Linux/macOS 课件常见打包方式）：与 zip 走同一条抽取链路。"""
    import tarfile

    package = io.BytesIO()
    with tarfile.open(fileobj=package, mode="w:gz") as tf:
        for name, text in (("课件/大纲.txt", "第 1-8 周：进程与内存管理"),
                           ("课件/作业.csv", "周次,题目\n3,分页置换\n")):
            data = text.encode()
            info = tarfile.TarInfo(name)
            info.size = len(data)
            tf.addfile(info, io.BytesIO(data))

    body = body_of(client, "课件.tar.gz", package.getvalue())
    assert body["format"]["ext"] == ".tar.gz"
    assert body["format"]["group"] == "archive"
    preview = body["extracted_text_preview"]
    assert "进程与内存管理" in preview
    assert "分页置换" in preview


def test_upload_single_gzip_stream(client):
    """单文件压缩流（讲义.txt.gz）：解压后按真实内容抽取，不是「压缩包损坏」。"""
    import gzip

    body = body_of(client, "讲义.txt.gz", gzip.compress("第三章 进程调度与死锁".encode()))
    assert body["format"]["group"] == "archive"
    assert "进程调度" in body["extracted_text_preview"]


def test_upload_7z_gives_actionable_guidance(client):
    """7z/rar 本机没有解码器：回包必须说清「先用工具解压再上传」，而不是认不出。"""
    body = body_of(client, "资料.7z", b"7z\xbc\xaf\x27\x1c" + b"\x00" * 128)
    assert body["format"]["ext"] == ".7z"
    assert body["format"]["supported"] is False
    assert body["format"]["note"]                       # 前端灰显时的悬停说明
    warnings = " ".join(body["extract"]["warnings"])
    assert "解压" in warnings
    assert "未识别" not in warnings


# ---------------------------------------------------------------------------
# 嗅探纠正 / 不支持类型的引导
# ---------------------------------------------------------------------------

def test_upload_extension_mismatch_is_reported_and_corrected(client):
    """把 PDF 命名成 .pptx：按真实内容解析，并提示扩展名不一致。"""
    pdf = _pdf_bytes("Operating Systems homework. The deadline is due 2026-12-01 23:59.")
    body = body_of(client, "其实是PDF.pptx", pdf)
    assert body["format"]["ext"] == ".pdf"
    assert body["format"]["mismatch"] is True
    assert any("不一致" in w for w in body["extract"]["warnings"])


def test_upload_pdf_txt_extension(client):
    """把 PDF 命名成 .txt 也要能识别（内容优先）。"""
    body = body_of(client, "作业.txt", _pdf_bytes("Submit the report before 2026-12-02 23:59."))
    assert body["format"]["ext"] == ".pdf"


def test_upload_media_gives_actionable_hint(client):
    """音视频：明确说「不支持离线转写」并给出替代做法，而不是「未知格式」。"""
    body = body_of(client, "录音.mp3", b"ID3\x03\x00\x00\x00" + b"\x00" * 512)
    assert body["format"]["group"] == "media"
    assert body["format"]["supported"] is False
    warnings = " ".join(body["extract"]["warnings"])
    assert "转写" in warnings and "课件" in warnings


def test_upload_unknown_binary(client):
    """完全认不出的类型：给出「转成 PDF/TXT」的指引，不报 500。"""
    body = body_of(client, "神秘文件.zzz", b"\x00\x01\x02\x03\x04binary")
    assert body["format"]["group"] == "other"
    warnings = " ".join(body["extract"]["warnings"])
    assert "未能识别" in warnings


def test_upload_damaged_docx(client):
    """损坏的 docx：抽取失败但接口仍 200，且给出中文提示。"""
    body = body_of(client, "坏文件.docx", b"PK\x03\x04" + b"\x00" * 100)
    assert body["text_chars"] == 0
    assert body["extract"]["warnings"]


def test_upload_image_without_ocr_keeps_guidance(client):
    """OCR 不可用（测试环境强制关闭）时，图片上传给出可操作提示且不崩溃。"""
    png = _png_bytes()
    body = body_of(client, "板书截图.png", png, "image/png")
    assert body["format"]["group"] == "image"
    assert body["format"]["ocr"] is True
    warnings = " ".join(body["extract"]["warnings"])
    assert "OCR" in warnings or "识别" in warnings


def test_upload_heic_without_decoder(client):
    """HEIC 缺解码器：明确提示先转 JPG/PNG，而不是「文件已损坏」。"""
    from app.infra.ingest.formats import decoder_available

    if decoder_available(".heic"):
        pytest.skip("本机有 HEIC 解码器")
    data = b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00heicmif1" + b"\x00" * 64
    r = upload(client, "手机照片.heic", data, "image/heic")
    assert r.status_code == 422
    assert "JPG" in r.json()["detail"] or "PNG" in r.json()["detail"]


def test_format_catalog_endpoint(client):
    """`GET /api/ingest/formats`：前端 accept 与格式矩阵的数据源。"""
    r = client.get("/api/ingest/formats")
    assert r.status_code == 200
    data = r.json()
    assert data["groups"]
    assert ".pdf" in data["accept"].split(",")
    assert data["max_upload_mb"] > 0
    exts = {i["ext"] for g in data["groups"] for i in g["formats"]}
    assert {".pdf", ".docx", ".xlsx", ".pptx", ".epub", ".ipynb", ".srt", ".zip", ".png"} <= exts
    # 压缩包覆盖：zip / tar 家族 / 单文件压缩流 / 缺解码器的 7z·rar 都在目录里
    assert {".tar", ".tar.gz", ".tgz", ".gz", ".bz2", ".xz", ".7z", ".rar"} <= exts
    archive = [g for g in data["groups"] if g["id"] == "archive"][0]
    supported = {i["ext"]: i["supported"] for i in archive["formats"]}
    assert supported[".zip"] is True and supported[".tar.gz"] is True
    assert supported[".7z"] is False and supported[".rar"] is False
    assert all(i["label"] for g in data["groups"] for i in g["formats"])


def test_uploaded_file_lands_in_store(client):
    """上传的文件必须落在 data/files 内（确认入库时按路径回填）。"""
    from app.infra.filestore import files_root

    body = body_of(client, "落盘检查.md", "# 标题\n正文".encode("utf-8"))
    stored = Path(body["stored_path"])
    assert stored.exists()
    assert stored.resolve().is_relative_to(files_root().resolve())


# ---------------------------------------------------------------------------
# fixture 构造
# ---------------------------------------------------------------------------

def _pdf_bytes(text: str) -> bytes:
    """用 PyMuPDF 生成单页 PDF（与用户上传的真实文件同一解析路径）。

    必须经 `office._pymupdf()` 导入：PyMuPDF 的 SWIG 扩展在导入期发
    DeprecationWarning，而 pytest.ini 把该告警升级为错误 —— 裸 import 会让
    异常从 C 扩展初始化里抛出，直接把整个测试进程打成段错误。
    """
    from app.infra.ingest.office import _pymupdf

    pymupdf = _pymupdf()
    if pymupdf is None:  # pragma: no cover - 缺依赖时跳过
        pytest.skip("PyMuPDF unavailable")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), text)
    data = doc.tobytes()
    doc.close()
    return data


def _png_bytes() -> bytes:
    return __import__("base64").b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
    )


def _odt_bytes(body_xml: str) -> bytes:
    """最小 ODT（content.xml 里放给定正文片段）。"""
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
        f"<office:body><office:text>{body_xml}</office:text></office:body>"
        "</office:document-content>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr("META-INF/manifest.xml", "<manifest/>")
        zf.writestr("content.xml", content)
    return buf.getvalue()


def _epub_bytes(chapters: list[tuple[str, str]]) -> bytes:
    """最小 EPUB（mimetype + container.xml + OPF spine + 章节）。"""
    items = "".join(
        f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>'
        for i, (name, _) in enumerate(chapters)
    )
    spine = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
    opf = (
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        f"<manifest>{items}</manifest><spine>{spine}</spine></package>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        zf.writestr("OEBPS/content.opf", opf)
        for name, html in chapters:
            zf.writestr(f"OEBPS/{name}", html)
    return buf.getvalue()
